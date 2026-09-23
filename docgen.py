import io
from pathlib import Path
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from table_helpers import (
    set_font, add_question_heading, add_question_table, add_question_total_row,
    add_grand_total_paragraph, add_marking_guide_question_table,
    set_document_page_setup, set_document_default_font,
    add_information_sheet_image, set_table_fixed_layout_and_grid,
    show_table_borders
)

DEFAULT_COGNITIVE_TARGETS = {
    'knowledge': 20,
    'routine': 35,
    'complex': 30,
    'problem_solving': 15
}

COGNITIVE_LEVELS = ['knowledge', 'routine', 'complex', 'problem_solving']
COGNITIVE_GRID_COL_WIDTHS_IN = [1.3, 0.85, 0.85, 0.85, 1.15, 0.7]


class DocumentGenerator:
    def __init__(self, template_path: str = None):
        """
        Args:
            template_path: Path to a .docx template with {{PLACEHOLDER}} tokens.
                          If None, create a generic cover page (per PROJECT_BRIEF.md:
                          plain, no branding, no border image -- this wins over
                          task_SKILL.md's richer branded front matter for the MVP).
        """
        self.template_path = template_path
        self.info_sheet_path = Path(__file__).parent / "information_sheet_gr11_gr12.png"

    def generate_full_assessment(self, topics_data: list,
                                total_marks: int, task: str = None,
                                term: str = None, time_minutes: int = None,
                                examiner: str = None, moderator: str = None,
                                grade: str = "11",
                                target_distribution: dict = None,
                                include_information_sheet: bool = True) -> bytes:
        """
        Generate the complete assessment as ONE .docx: question paper, marking
        guide, and cognitive level analysis grid, followed by the information
        sheet (Grade 11/12). This matches the proven Task 7 precedent
        (a single combined _QP_MG.docx) and PROJECT_BRIEF.md's framing of the
        product as one downloadable document per generated assessment.

        Args:
            topics_data: List of (topic_name: str, questions: list[GeneratedQuestion])
                       tuples, in the order they should appear. Each topic becomes
                       its own QUESTION N group (QUESTION 1, QUESTION 2, ...), each
                       internally sub-numbered N.1, N.2... -- matching how a real
                       multi-topic DBE paper structures one Question block per topic.
                       Single-topic papers just pass a one-item list.
            total_marks: Total marks across all topics combined
            task, term, time_minutes, examiner, moderator: Metadata for cover page/header
            grade: Grade (default "11")
            target_distribution: Cognitive % targets for the FET TARGET row (defaults
                                  to CAPS balanced K20/R35/C30/PS15)
            include_information_sheet: Append the DBE info sheet as the last page
                                        (Grade 11/12 only; per task_SKILL.md)

        Returns:
            .docx file as bytes
        """
        target_distribution = target_distribution or DEFAULT_COGNITIVE_TARGETS
        topic_display = " & ".join(name for name, _ in topics_data)

        if self.template_path:
            doc = self._load_and_replace_template(
                topic_display, task, term, time_minutes, total_marks, examiner,
                moderator, grade
            )
            # Per format_SKILL.md: a real template supplies its own header/footer/
            # border via its own section properties -- don't override them.
        else:
            doc = Document()
            set_document_page_setup(doc)
            set_document_default_font(doc)
            self._add_generic_cover_page(doc, topic_display, task, term, time_minutes,
                                         total_marks, examiner, moderator, grade)
            # Per today's ruling: header text + page-number footer apply even
            # in the no-template path (task_SKILL.md's frame elements, applied
            # regardless of PROJECT_BRIEF.md's plain-cover-page simplification).
            left_header = f"Mathematics {task}".strip() if task else f"Grade {grade} Mathematics"
            right_header = term or ""
            self._set_header_footer(doc, left_header, right_header)

        # Each selected topic becomes its own QUESTION N group, sub-numbered
        # N.1, N.2, N.3... (see audit note: real DBE papers sometimes group
        # same-flavour "solve for x" items under a shared N.1 stem with
        # N.1.1/N.1.2 children within a topic -- that finer stem-grouping
        # heuristic is a follow-up enhancement, not implemented here; flat
        # depth-0 sub-numbering under each QUESTION heading is fully
        # spec-compliant per format_SKILL.md's own column-merge and
        # blank-spacer-row rules).
        question_groups = [(i + 1, questions) for i, (_, questions) in enumerate(topics_data)]

        self._add_question_paper_body(doc, question_groups, total_marks)
        doc.add_page_break()
        self._add_marking_guide_body(doc, question_groups, total_marks)
        self._add_cognitive_grid(doc, question_groups, total_marks, target_distribution)

        if include_information_sheet and grade in ("11", "12"):
            doc.add_page_break()
            add_information_sheet_image(doc, str(self.info_sheet_path))

        return self._save_to_bytes(doc)

    def _load_and_replace_template(self, topic: str, task: str, term: str,
                                   time_minutes: int, total_marks: int,
                                   examiner: str, moderator: str, grade: str):
        """Load a template and replace {{PLACEHOLDER}} tokens."""
        doc = Document(self.template_path)

        replacements = {
            "{{GRADE}}": grade,
            "{{TASK}}": task or "",
            "{{TERM}}": term or "",
            "{{TOTAL}}": str(total_marks),
            "{{TIME}}": f"{time_minutes} minutes" if time_minutes else "",
            "{{EXAMINER}}": examiner or "",
            "{{MODERATOR}}": moderator or "",
            "{{TOPIC}}": topic
        }

        for paragraph in doc.paragraphs:
            for key, value in replacements.items():
                if key in paragraph.text:
                    for run in paragraph.runs:
                        if key in run.text:
                            run.text = run.text.replace(key, value)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for key, value in replacements.items():
                            if key in paragraph.text:
                                for run in paragraph.runs:
                                    if key in run.text:
                                        run.text = run.text.replace(key, value)

        return doc

    def _add_generic_cover_page(self, doc, topic: str, task: str, term: str,
                                time_minutes: int, total_marks: int,
                                examiner: str, moderator: str, grade: str):
        """Plain generic cover page -- no branding, no border image, no logo,
        per PROJECT_BRIEF.md's explicit MVP instruction (wins over task_SKILL.md's
        richer branded two-page front matter for this no-template path)."""
        title = doc.add_paragraph()
        r = title.add_run(f"Grade {grade} Mathematics")
        set_font(r, size=16, bold=True)

        subtitle = doc.add_paragraph()
        r = subtitle.add_run(topic)
        set_font(r, size=14, bold=True)

        info_items = []
        if task:
            info_items.append(f"Task: {task}")
        if term:
            info_items.append(f"Term: {term}")
        if time_minutes:
            info_items.append(f"Time: {time_minutes} minutes")
        info_items.append(f"Total Marks: {total_marks}")
        if examiner:
            info_items.append(f"Examiner: {examiner}")
        if moderator:
            info_items.append(f"Moderator: {moderator}")

        for item in info_items:
            p = doc.add_paragraph()
            r = p.add_run(item)
            set_font(r, size=12)

        doc.add_page_break()

    def _set_header_footer(self, doc, left_text: str, right_text: str):
        """Left/right header text + right-aligned page-number footer field.
        Applied to the no-template generic-cover path per today's ruling."""
        section = doc.sections[0]

        header = section.header
        header.is_linked_to_previous = False
        p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        for run in list(p.runs):
            run.text = ''
        usable_width = section.page_width - section.left_margin - section.right_margin
        p.paragraph_format.tab_stops.add_tab_stop(usable_width, WD_TAB_ALIGNMENT.RIGHT)
        r_left = p.add_run(left_text)
        set_font(r_left, size=10)
        p.add_run('\t')
        r_right = p.add_run(right_text)
        set_font(r_right, size=10)

        footer = section.footer
        footer.is_linked_to_previous = False
        fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        for run in list(fp.runs):
            run.text = ''
        fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        self._add_page_number_field(fp)

    def _add_page_number_field(self, paragraph):
        """Insert a real Word PAGE field (updates automatically), not hand-typed text."""
        run = paragraph.add_run()
        set_font(run, size=10)
        fld_begin = OxmlElement('w:fldChar')
        fld_begin.set(qn('w:fldCharType'), 'begin')
        instr = OxmlElement('w:instrText')
        instr.set(qn('xml:space'), 'preserve')
        instr.text = 'PAGE'
        fld_end = OxmlElement('w:fldChar')
        fld_end.set(qn('w:fldCharType'), 'end')
        run._r.append(fld_begin)
        run._r.append(instr)
        run._r.append(fld_end)

    def _add_question_paper_body(self, doc, question_groups: list, grand_total: int):
        """One QUESTION N heading + one 4-column hidden-border table per group,
        sub-questions numbered N.1, N.2..., a per-question [N] total row inside
        the table, then the standalone bold centred grand TOTAL [N] paragraph."""
        for qnum, questions in question_groups:
            add_question_heading(doc, f"QUESTION {qnum}")
            rows = [
                {'num': f'{qnum}.{i}', 'parts': [q.question_text], 'marks': q.marks}
                for i, q in enumerate(questions, 1)
            ]
            table = add_question_table(doc, rows)
            question_total = sum(q.marks for q in questions)
            add_question_total_row(table, question_total)
            doc.add_paragraph()

        add_grand_total_paragraph(doc, grand_total)

    def _add_marking_guide_body(self, doc, question_groups: list, grand_total: int):
        """3-column visible-border marking guide table, same N.1/N.2 numbering
        as the question paper, one calculation step per row, ending with the
        same standalone grand TOTAL [N] paragraph."""
        heading = doc.add_paragraph()
        r = heading.add_run("MARKING GUIDELINE")
        set_font(r, size=14, bold=True)
        doc.add_paragraph()

        for qnum, questions in question_groups:
            step_groups = []
            for i, q in enumerate(questions, 1):
                steps = q.marking_steps or [{"text": q.answer_text, "tick_label": "answer", "tick_count": 1}]
                step_tuples = []
                for step in steps:
                    count = step.get("tick_count", 0)
                    label = step.get("tick_label")
                    tick_display = f"{'✓' * count} {label}" if count > 0 and label else None
                    step_tuples.append(([step.get("text", "")], tick_display))
                step_groups.append({
                    'num': f'{qnum}.{i}',
                    'steps': step_tuples,
                    'marks': q.marks
                })
            add_marking_guide_question_table(doc, qnum, step_groups)
            doc.add_paragraph()

        add_grand_total_paragraph(doc, grand_total)

    def _add_cognitive_grid(self, doc, question_groups: list, total_marks: int, targets: dict):
        """Bordered COGNITIVE LEVEL ANALYSIS GRID -- long format by question,
        with FET TARGET % and THIS TASK % rows, per task_SKILL.md."""
        doc.add_paragraph()
        heading = doc.add_paragraph()
        r = heading.add_run("COGNITIVE LEVEL ANALYSIS GRID")
        set_font(r, size=14, bold=True)

        table = doc.add_table(rows=0, cols=6)
        show_table_borders(table)

        headers = ["QUESTION", "Knowledge", "Routine", "Complex", "Problem-Solving", "Total"]
        header_row = table.add_row()
        for i, h in enumerate(headers):
            r = header_row.cells[i].paragraphs[0].add_run(h)
            set_font(r, size=11, bold=True)

        level_totals = {level: 0 for level in COGNITIVE_LEVELS}

        for qnum, questions in question_groups:
            row_totals = {level: 0 for level in COGNITIVE_LEVELS}
            for q in questions:
                if q.cognitive_level in row_totals:
                    row_totals[q.cognitive_level] += q.marks
            row_total_marks = sum(row_totals.values())

            row = table.add_row()
            r = row.cells[0].paragraphs[0].add_run(f"QUESTION {qnum}")
            set_font(r, size=11)
            for i, level in enumerate(COGNITIVE_LEVELS):
                r = row.cells[i + 1].paragraphs[0].add_run(str(row_totals[level]))
                set_font(r, size=11)
                level_totals[level] += row_totals[level]
            r = row.cells[5].paragraphs[0].add_run(str(row_total_marks))
            set_font(r, size=11)

        total_row = table.add_row()
        r = total_row.cells[0].paragraphs[0].add_run("TOTAL")
        set_font(r, size=11, bold=True)
        for i, level in enumerate(COGNITIVE_LEVELS):
            r = total_row.cells[i + 1].paragraphs[0].add_run(str(level_totals[level]))
            set_font(r, size=11, bold=True)
        r = total_row.cells[5].paragraphs[0].add_run(str(sum(level_totals.values())))
        set_font(r, size=11, bold=True)

        fet_row = table.add_row()
        r = fet_row.cells[0].paragraphs[0].add_run("FET TARGET %")
        set_font(r, size=11, bold=True)
        for i, level in enumerate(COGNITIVE_LEVELS):
            r = fet_row.cells[i + 1].paragraphs[0].add_run(f"{targets.get(level, 0)}%")
            set_font(r, size=11, bold=True)
        r = fet_row.cells[5].paragraphs[0].add_run("—")
        set_font(r, size=11, bold=True)

        task_row = table.add_row()
        r = task_row.cells[0].paragraphs[0].add_run("THIS TASK %")
        set_font(r, size=11, bold=True)
        for i, level in enumerate(COGNITIVE_LEVELS):
            pct = round(level_totals[level] / total_marks * 100) if total_marks else 0
            r = task_row.cells[i + 1].paragraphs[0].add_run(f"{pct}%")
            set_font(r, size=11, bold=True)
        r = task_row.cells[5].paragraphs[0].add_run("100%")
        set_font(r, size=11, bold=True)

        set_table_fixed_layout_and_grid(table, COGNITIVE_GRID_COL_WIDTHS_IN)

    def _save_to_bytes(self, doc) -> bytes:
        """Save document to bytes buffer."""
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()
